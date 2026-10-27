#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
check_source.py — 送檔前的「來源檔完整性自檢」（在使用者 Windows 筆電執行）

用途
  RefreshAll 有時會跑不完（列數最多的「2610WEB&KOL」樞紐表最常被切掉），
  造成某些分頁停留在前一天。這種缺漏「不會報錯」，數字看起來也合理，
  推上儀表板後會變成「某天業績崩盤」的假象。
  2026-10-19／20／21 連續三天踩到，每次都要重跑、重傳，來回 10〜20 分鐘。
  本腳本在送檔前 5 秒內把這件事擋下來。

用法（擇一）
  1) 雙擊同資料夾的  check_source.bat
  2) 命令列：  python check_source.py
     指定檔案：python check_source.py "D:\\path\\FY26商售實績明細(claude).xlsx"
     指定日期：python check_source.py --date 2026-10-20

檢查內容（當月四個分頁）
  · 最後日期：WEB&KOL／EC／WEB數位典藏 要到期望日；TM（電銷）週末與國定假日不上班，
    只要到「期望日當天或之前的最後一個工作日」（工作日＝週一〜五、且不在 data.json 的 holidays）
  · 中間缺日（月初〜期望日）：WEB&KOL、EC 任何一天 0 列；TM 任何工作日 0 列；
    WEB數位典藏一天只有 1〜2 筆，0 列也列出，但可能是真的沒業績，請人工確認
  · 使用者確認過「真的沒資料」的日子記在 check_source_confirmed.json（日期 → 分頁），不擋
    （2026-10-27 起：連假後 TM 10/25〜26 被誤判落後，而 TM 10/23 的 0 列反而看不到）

預設行為
  · 未指定檔案時，讀 daily_run.local.json 的 master_xlsx
  · 未指定日期時，期望值＝「昨天」
  · 回傳碼 0＝全部通過、1＝有分頁沒跟上或有 0 列的日子（需要人看）、2＝設定或檔案問題

相依套件
  pip install openpyxl
"""

import sys, os, json, datetime, collections, argparse, unicodedata

HERE = os.path.dirname(os.path.abspath(__file__))
CONFIG = os.path.join(HERE, 'daily_run.local.json')
CONFIRMED = os.path.join(HERE, 'check_source_confirmed.json')
DATA_JSON = os.path.join(os.path.dirname(HERE), 'data.json')

# 每日實績檔裡「當月」會用到的分頁字首（月份字首在執行時動態組出來）
SHEET_SUFFIXES = ['WEB&KOL', 'EC', 'TM', 'WEB數位典藏']
WORKDAY_ONLY = {'TM'}             # 只在工作日有資料
MAYBE_ZERO = {'WEB數位典藏'}       # 0 列可能是真的沒業績
WEEKDAY = '一二三四五六日'

GREEN = '\033[92m'; RED = '\033[91m'; YEL = '\033[93m'; OFF = '\033[0m'
if os.name == 'nt':
    try:
        import colorama; colorama.just_fix_windows_console()
    except Exception:
        GREEN = RED = YEL = OFF = ''


def parse_date(v):
    """Excel 的日期可能是 datetime、date，或 '2026/10/1'、'2026-10-01 00:00:00' 這類字串。"""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        p = v.strip()[:10].replace('/', '-').split('-')
        if len(p) == 3:
            try:
                return datetime.date(int(p[0]), int(p[1]), int(p[2]))
            except ValueError:
                return None
    return None


def fy_month_prefix(d):
    """2026-10-20 → '2610'（分頁字首）。"""
    return '%02d%02d' % (d.year % 100, d.month)


def pad(s, w=22):
    """中文字佔兩格，補空白對齊。"""
    n = sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in s)
    return s + ' ' * max(0, w - n)


def fmt_days(days):
    """[2026-10-23, ...] → '10/23（五）、...'"""
    return '、'.join(f'{d.month}/{d.day}（{WEEKDAY[d.weekday()]}）' for d in days)


def is_workday(d, holidays):
    return d.weekday() < 5 and d not in holidays


def last_workday(d, holidays):
    while not is_workday(d, holidays):
        d -= datetime.timedelta(days=1)
    return d


def load_holidays():
    with open(DATA_JSON, encoding='utf-8') as f:
        hol = json.load(f).get('holidays')
    if not isinstance(hol, list):
        raise ValueError('data.json 沒有 holidays 清單')
    return {d for d in map(parse_date, hol) if d}


def load_confirmed():
    """{"YYYY-MM-DD": ["TM", ...]} → {date: {'TM', ...}}；底線開頭的鍵是說明。檔案不存在＝空清單。"""
    if not os.path.exists(CONFIRMED):
        return {}
    with open(CONFIRMED, encoding='utf-8') as f:
        raw = json.load(f)
    out = {}
    for k, v in raw.items():
        if k.startswith('_'):
            continue
        d = parse_date(k)
        if d is None or not isinstance(v, list) or not set(v) <= set(SHEET_SUFFIXES):
            raise ValueError(f'格式不對：{k!r}: {v!r}（分頁只能是 {"／".join(SHEET_SUFFIXES)}）')
        out[d] = set(v)
    return out


def sheet_last_date(ws):
    """回傳 (總列數, 最後日期, {日期: 列數})；有日期欄但沒有資料時回 (0, None, {})。找不到日期欄時回 None。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, r in enumerate(rows[:8]):
        if r and any(isinstance(x, str) and '日期' in x for x in r if x):
            hdr = i
            break
    if hdr is None:
        return None
    j = next(k for k, x in enumerate(rows[hdr]) if isinstance(x, str) and '日期' in x)
    per = collections.Counter()
    for r in rows[hdr + 1:]:
        if j < len(r):
            d = parse_date(r[j])
            if d:
                per[d] += 1
    if not per:
        return 0, None, {}
    return sum(per.values()), max(per), per


def main():
    ap = argparse.ArgumentParser(description='每日實績檔的分頁完整性自檢')
    ap.add_argument('xlsx', nargs='?', help='實績明細檔路徑（省略則讀 daily_run.local.json）')
    ap.add_argument('--date', help='期望的最後日期 YYYY-MM-DD（省略＝昨天）')
    args = ap.parse_args()

    # ── 檔案路徑 ──
    path = args.xlsx
    if not path:
        if not os.path.exists(CONFIG):
            print(f'{RED}找不到設定檔 {CONFIG}{OFF}')
            print('請把檔案路徑當參數傳入，或先建立 daily_run.local.json。')
            return 2
        with open(CONFIG, encoding='utf-8') as f:
            path = json.load(f).get('master_xlsx')
    if not path or not os.path.exists(path):
        print(f'{RED}找不到實績檔：{path}{OFF}')
        return 2

    # ── 假日與確認清單 ──
    try:
        holidays = load_holidays()
    except (OSError, ValueError) as e:
        print(f'{RED}讀不到 {DATA_JSON} 的 holidays：{e}{OFF}')
        return 2
    try:
        confirmed = load_confirmed()
    except (OSError, ValueError) as e:
        print(f'{RED}確認清單 {os.path.basename(CONFIRMED)} 讀不進來：{e}{OFF}')
        return 2

    # ── 期望日期 ──
    if args.date:
        want = parse_date(args.date)
        if want is None:
            print(f'{RED}--date 格式不對：{args.date}（要 YYYY-MM-DD）{OFF}')
            return 2
    else:
        want = datetime.date.today() - datetime.timedelta(days=1)
    prefix = fy_month_prefix(want)
    month_start = want.replace(day=1)
    want_tm = last_workday(want, holidays)

    mtime = datetime.datetime.fromtimestamp(os.path.getmtime(path))
    print(f'檔案：{os.path.basename(path)}')
    print(f'存檔時間：{mtime:%Y-%m-%d %H:%M}')
    print(f'期望最後日期：{want}（{prefix} 分頁）；TM 只看工作日 → {want_tm}')
    print(f'缺日檢查：{month_start}〜期望日，0 列的日子都列出（確認清單上的不擋）')
    print('-' * 66)

    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    found = {}
    for sn in wb.sheetnames:
        s = sn.strip()
        if not s.startswith(prefix):
            continue
        found[s] = sheet_last_date(wb[sn])

    rerun, ask, bad, missing_sheets = [], [], [], []
    for suf in SHEET_SUFFIXES:
        name = prefix + suf
        want_s = want_tm if suf in WORKDAY_ONLY else want
        if want_s < month_start:
            print(f'{GREEN}✓ {pad(name)} 本月還沒有工作日，不檢查{OFF}')
            continue
        hit = next((k for k in found if k.replace(' ', '') == name.replace(' ', '')), None)
        if hit is None:
            print(f'{YEL}？ {pad(name)} 找不到這個分頁{OFF}')
            missing_sheets.append(name)
            continue
        if found[hit] is None:
            print(f'{RED}✗ {pad(hit)} 讀不到日期欄{OFF}')
            bad.append(hit)
            continue
        n, mx, per = found[hit]

        days = [month_start + datetime.timedelta(days=i) for i in range((want_s - month_start).days + 1)]
        if suf in WORKDAY_ONLY:
            days = [d for d in days if is_workday(d, holidays)]
        zero = [d for d in days if not per.get(d)]
        ok_zero = [d for d in zero if suf in confirmed.get(d, ())]
        tail = [d for d in zero if d not in ok_zero and (mx is None or d > mx)]
        holes = [d for d in zero if d not in ok_zero and mx is not None and d < mx]

        if not tail:
            due = f'當日 {per.get(want, 0)} 列' if want_s == want else f'期望 {want_s}（最後一個工作日）'
            mark = f'{RED}✗' if holes else f'{GREEN}✓'
            print(f'{mark} {pad(hit)} 列數 {n:>5}　最後 {mx or "—"}　{due}{OFF}')
        elif suf in MAYBE_ZERO:
            print(f'{RED}✗ {pad(hit)} 列數 {n:>5}　最後 {mx or "—"}　{fmt_days(tail)} 0 列 ← 可能是真的沒業績，請人工確認{OFF}')
            ask.append(hit)
        elif mx is None:
            print(f'{RED}✗ {pad(hit)} 本月還沒有任何資料 ← 請重跑這個分頁{OFF}')
            rerun.append(hit)
        else:
            print(f'{RED}✗ {pad(hit)} 列數 {n:>5}　最後 {mx}　落後 {(want_s - mx).days} 天 ← 請重跑這個分頁{OFF}')
            rerun.append(hit)
        if holes:
            note = '可能是真的沒業績，請人工確認' if suf in MAYBE_ZERO else '請確認是沒刷到還是真的沒資料'
            print(f'{RED}    ✗ 中間 0 列：{fmt_days(holes)} ← {note}{OFF}')
            if hit not in ask:
                ask.append(hit)
        if ok_zero:
            print(f'    · 已確認無資料：{fmt_days(ok_zero)}')
    wb.close()

    print('-' * 66)
    if not (rerun or ask or bad or missing_sheets):
        print(f'{GREEN}全部分頁都到期望日、本月沒有未確認的缺日，可以送出。{OFF}')
        return 0
    if missing_sheets:
        print(f'{YEL}有分頁找不到（可能是分頁名稱改了），請人工確認。{OFF}')
    if rerun or bad:
        print(f'{RED}有分頁沒跟上。請在 Excel 重跑 RefreshAll（或對該分頁的樞紐表按右鍵→重新整理），{OFF}')
        print(f'{RED}存檔後再執行一次本檢查。{OFF}')
    if ask:
        print(f'{RED}有 0 列的日子（{"、".join(ask)}）。沒刷到就在 Excel 重跑；{OFF}')
        print(f'{RED}確認是真的沒資料，就把日期與分頁加進 scripts/{os.path.basename(CONFIRMED)} 後再跑一次。{OFF}')
    return 1


if __name__ == '__main__':
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        sys.exit(130)
