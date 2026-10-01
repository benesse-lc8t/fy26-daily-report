#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把「巧虎雙語點讀圖鑑套組預購統計」Excel 寫入 data.json 的 preorderLedger。

用法：
    python3 scripts/update_preorder_ledger.py <套組預購統計.xlsx> [--rule dianduo-tujian]

── 為什麼需要這支腳本 ──────────────────────────────────────────
每日實績 Excel 與這份預購台帳的涵蓋範圍**不同**，兩邊都不完整：

  · EC / WEB / KOL 以「訂單日」開列 —— 尚未出貨就會進每日明細，
    `update_data.py` 的 PREORDER_RULES 會自動排除，無須人工標註。
  · TM 以「發貨日」開列 —— 預購單在「未發送」狀態下根本不會進每日明細，
    所以每日 Excel 看不到 TM 的預購接單量。

因此：
  每日 Excel  → 「已進銷售明細、已排除」（分母是會影響實績的部分）
  本預購台帳  → 「預購總接單」（含尚未進銷售明細的 TM 未發送單）
兩者的差額本身就是有意義的資訊（通常＝TM 未發送量 ± 兩邊截止日落差）。

── 來源檔版面（請勿變動欄位位置）────────────────────────────────
分頁「套組預購數量&金額」
  · 第 2 列：統計期間、單套未稅金額
  · 第 3 列：再次上市日期
  · 標題列：含 WEB / KOL / 蝦皮 / TM 各兩組（前為套數、後為未稅金額）
  · 資料列：該列第一個 datetime 儲存格＝日期，其後依標題欄位取值
腳本以「標題列文字」定位欄位，容忍整體位移，但不容忍欄位被插入／刪除。

小計欄一律**由腳本重算**，不採用來源檔的小計（來源檔公式曾漏計 TM）。
"""
import json
import os
import sys
from datetime import datetime, timezone

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_JSON = os.path.join(ROOT, 'data.json')

SHEET = '套組預購數量&金額'
# 來源檔的通路名 → data.json 的通路代碼
CH_MAP = {'WEB': 'WEB', 'KOL': 'KOL', '蝦皮': 'EC', 'TM': 'TM'}
CH_ORDER = ['WEB', 'KOL', '蝦皮', 'TM']


def find_header(ws, max_scan=20, max_col=24):
    """找標題列，回傳 (列號, 套數欄 index dict, 金額欄 index dict)。"""
    for r, row in enumerate(ws.iter_rows(max_row=max_scan, max_col=max_col, values_only=True), 1):
        vals = [str(v).strip() if v is not None else '' for v in row]
        hits = [i for i, v in enumerate(vals) if v == 'WEB']
        if len(hits) < 2:
            continue
        qty_i, amt_i = {}, {}
        ok = True
        for base, target in ((hits[0], qty_i), (hits[1], amt_i)):
            for k, name in enumerate(CH_ORDER):
                j = base + k
                if j >= len(vals) or vals[j] != name:
                    ok = False
                    break
                target[name] = j
            if not ok:
                break
        if ok:
            return r, qty_i, amt_i
    raise SystemExit(f'找不到標題列（需含兩組 WEB/KOL/蝦皮/TM）：{SHEET}')


def cell_num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def main(path, rule_id='dianduo-tujian'):
    print(f'Loading {path} …')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f'找不到分頁「{SHEET}」。現有分頁：{wb.sheetnames}')
    ws = wb[SHEET]

    hr, qty_i, amt_i = find_header(ws)
    print(f'  標題列 r{hr}　套數欄 {qty_i}　金額欄 {amt_i}')

    # 期間／單價等 meta。標籤與值的相對位置不固定：
    # 「統計期間」「再次上市日期」的值在同列右方，「單套未稅金額」的值在正下方。
    # 故先找右方、再找下方，並排除「值本身又是另一個標籤」的情況。
    LABELS = ('統計期間', '再次上市日期', '單套未稅金額')
    head = [list(r) for r in ws.iter_rows(max_row=hr - 1, max_col=20, values_only=True)]
    meta = {}

    def norm(v):
        return v.strftime('%Y-%m-%d') if isinstance(v, datetime) else v

    for ri, row in enumerate(head):
        for i, v in enumerate(row):
            label = str(v).strip() if v is not None else ''
            if label not in LABELS or label in meta:
                continue
            for j in range(i + 1, min(i + 4, len(row))):          # 同列右方
                cand = row[j]
                if cand is not None and str(cand).strip() not in LABELS \
                        and not str(cand).startswith('FY26'):
                    meta[label] = norm(cand)
                    break
            if label in meta:
                continue
            for rj in range(ri + 1, min(ri + 3, len(head))):      # 正下方
                cand = head[rj][i] if i < len(head[rj]) else None
                if cand is not None and str(cand).strip() not in LABELS:
                    meta[label] = norm(cand)
                    break

    rows, by_ch_q, by_ch_a = [], {c: 0.0 for c in CH_ORDER}, {c: 0.0 for c in CH_ORDER}
    maxcol = max(list(qty_i.values()) + list(amt_i.values())) + 1
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=maxcol, values_only=True):
        d = next((v for v in row[:4] if isinstance(v, datetime)), None)
        if d is None:
            continue
        q = {c: cell_num(row[qty_i[c]]) for c in CH_ORDER}
        a = {c: cell_num(row[amt_i[c]]) for c in CH_ORDER}
        if sum(q.values()) == 0 and sum(a.values()) == 0:
            continue                     # 尚無接單的日子不留列
        iso = d.strftime('%Y-%m-%d')
        rows.append({
            'date': iso,
            'qty': {CH_MAP[c]: q[c] for c in CH_ORDER if q[c] or a[c]},
            'amt': {CH_MAP[c]: round(a[c], 2) for c in CH_ORDER if q[c] or a[c]},
            'qtyTotal': sum(q.values()),                 # 小計一律重算
            'amtTotal': round(sum(a.values()), 2),
        })
        for c in CH_ORDER:
            by_ch_q[c] += q[c]
            by_ch_a[c] += a[c]

    tq = sum(by_ch_q.values())
    ta = sum(by_ch_a.values())
    print(f'\n  接單日數 {len(rows)}　合計 {tq:.0f} 套 / {ta:,.0f} 元')
    for c in CH_ORDER:
        if by_ch_q[c] or by_ch_a[c]:
            print(f'    {c:4s} {by_ch_q[c]:5.0f} 套 {by_ch_a[c]:>10,.0f} 元')
    print(f'  meta: {meta}')

    with open(DATA_JSON) as f:
        data = json.load(f)

    # 與每日實績檔分流出來的 preorderRows 對帳
    pre = [r for r in data.get('preorderRows', []) if r.get('preorderId') == rule_id]
    daily_q = {}
    daily_a = {}
    for r in pre:
        daily_q[r['ch']] = daily_q.get(r['ch'], 0) + (r['qty'] or 0)
        daily_a[r['ch']] = daily_a.get(r['ch'], 0) + r['amt']
    print(f'\n  對帳：台帳 {tq:.0f} 套 vs 每日檔已排除 {sum(daily_q.values()):.0f} 套'
          f'　差 {tq - sum(daily_q.values()):+.0f} 套')
    for c in CH_ORDER:
        code = CH_MAP[c]
        dq = daily_q.get(code, 0)
        if by_ch_q[c] or dq:
            flag = '' if abs(by_ch_q[c] - dq) < 0.01 else '  ← 差異'
            print(f'    {c:4s} 台帳 {by_ch_q[c]:4.0f} / 每日檔 {dq:4.0f}{flag}')

    data['preorderLedger'] = {
        'ruleId': rule_id,
        'sourceFile': os.path.basename(path),
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'period': meta.get('統計期間'),
        'relaunchDate': meta.get('再次上市日期'),
        'unitPriceExTax': meta.get('單套未稅金額'),
        'byChannelQty': {CH_MAP[c]: by_ch_q[c] for c in CH_ORDER},
        'byChannelAmt': {CH_MAP[c]: round(by_ch_a[c], 2) for c in CH_ORDER},
        'totalQty': tq,
        'totalAmt': round(ta, 2),
        'rows': rows,
    }

    with open(DATA_JSON, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    print(f'\ndata.json updated（只寫 preorderLedger，其餘鍵未動）')


if __name__ == '__main__':
    args = sys.argv[1:]
    rid = 'dianduo-tujian'
    if '--rule' in args:
        i = args.index('--rule')
        rid = args[i + 1]
        del args[i:i + 2]
    if not args:
        print('Usage: python3 scripts/update_preorder_ledger.py <套組預購統計.xlsx> [--rule <id>]')
        sys.exit(1)
    main(args[0], rid)
